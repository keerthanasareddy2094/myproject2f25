# resume_manager.py
import json, os, re
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, Optional

from pypdf import PdfReader
from docx import Document


from langchain_ollama import ChatOllama
from langchain_core.prompts import ChatPromptTemplate

DATA_DIR = Path("data")
DATA_DIR.mkdir(parents=True, exist_ok=True)

OLLAMA_HOST = os.getenv("OLLAMA_HOST", "http://127.0.0.1:11434")
MODEL_NAME = os.getenv("MODEL_NAME", "qwen2.5:0.5b")

# ---------- file -> text ----------
def _read_pdf(b: bytes) -> str:
    out = []
    reader = PdfReader(BytesIO(b))
    for p in reader.pages:
        try:
            out.append(p.extract_text() or "")
        except Exception:
            continue
    return "\n".join(out)

def _read_docx(b: bytes) -> str:
    buf = BytesIO(b)
    doc = Document(buf)
    return "\n".join([p.text for p in doc.paragraphs])

def read_file_to_text(uploaded) -> str:
    name = (uploaded.name or "").lower()
    b = uploaded.getvalue()
    if name.endswith(".pdf"):
        return _read_pdf(b)
    if name.endswith(".docx"):
        return _read_docx(b)
    try:
        return b.decode("utf-8", errors="ignore")
    except Exception:
        return b.decode("latin-1", errors="ignore")

def _llm(model_temp=0.1, num_ctx=4096, num_pred=400) -> ChatOllama:
    return ChatOllama(
        base_url=OLLAMA_HOST,
        model=MODEL_NAME,
        temperature=model_temp,
        streaming=False,
        model_kwargs={"num_ctx": num_ctx, "num_predict": num_pred},
    )

# ---------- LLM: résumé -> JSON (no hardcoded parsing) ----------
def llm_structured_resume(resume_text: str) -> Dict[str, Any]:
    """
    Single LLM call that returns compact JSON only.
    Literal braces are escaped so ChatPromptTemplate does not misread placeholders.
    """
    if not (resume_text or "").strip():
        return {}

    system = (
        "You are a precise résumé parser. Return ONLY compact JSON (no prose). "
        "Schema (omit empty fields): "
        "{{{{"
        '  "name": "string",'
        '  "email": "string",'
        '  "phone": "string",'
        '  "links": {{"linkedin":"url","github":"url","portfolio":"url","other":["url",...]}},'
        '  "summary": "1-2 sentences",'
        '  "skills": ["token",...],'
        '  "education": [{{"school":"","degree":"","field":"","start":"","end":"","gpa":""}}],'
        '  "experience": [{{"company":"","title":"","start":"","end":"","location":"","bullets":["..."]}}],'
        '  "projects": [{{"name":"","tech":["..."],"summary":""}}],'
        '  "certifications": ["..."]'
        "}}}}"
        "Rules: Be faithful to input text. Do not invent data. Lowercase skill tokens."
    )
    prompt = ChatPromptTemplate.from_messages([
        ("system", system),
        ("human", "RESUME TEXT:\n{resume_text}\n\nReturn JSON now.")
    ])
    resp = (prompt | _llm()).invoke({"resume_text": resume_text[:20000]})
    raw = resp.content or "{}"
    m = re.search(r"\{[\s\S]*\}", raw)
    json_str = m.group(0) if m else raw
    try:
        return json.loads(json_str)
    except Exception:
        return {}

# ---------- LLM: router (is the user asking about résumé?) ----------
def llm_is_resume_question(user_text: str) -> bool:
    system = (
        "Return JSON only. Decide if the user is asking about the résumé on file "
        'with yes/no: {{"resume_q": true|false}}. '
        "Treat queries like 'my name in resume', 'list my skills', 'show projects', "
        "'what is my linkedin', 'education from my cv' as true."
    )
    prompt = ChatPromptTemplate.from_messages([
        ("system", system),
        ("human", "{q}")
    ])
    out = (prompt | _llm(model_temp=0.0, num_ctx=512, num_pred=20)).invoke({"q": user_text}).content
    m = re.search(r"\{[\s\S]*\}", out or "")
    try:
        d = json.loads(m.group(0) if m else out)
        return bool(d.get("resume_q"))
    except Exception:
        return False

# ---------- LLM: grounded résumé QA (no handcoded mapping) ----------
def llm_answer_from_resume(user_text: str, resume_text: str, resume_json: Optional[Dict[str, Any]] = None) -> str:
    system = (
        "You are a concise assistant that answers ONLY using the provided résumé content. "
        "If the answer is not present, say you cannot find it in the résumé. "
        "Prefer exact values from JSON; otherwise quote short snippets from TEXT. "
        "Never fabricate."
    )
    prompt = ChatPromptTemplate.from_messages([
        ("system", system),
        ("human",
         "QUESTION:\n{q}\n\nRESUME JSON (may be partial):\n{j}\n\nRESUME TEXT:\n{t}\n\n"
         "Answer succinctly. If listing skills/education/experience, use short bullet points.")
    ])
    j = json.dumps(resume_json or {}, ensure_ascii=False)
    resp = (prompt | _llm(model_temp=0.1, num_ctx=4096, num_pred=280)).invoke({"q": user_text, "j": j, "t": resume_text[:20000]})
    return (resp.content or "").strip()

# ---------- persistence ----------
def save_resume(resume_text: str, resume_json: Dict[str, Any]) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    (DATA_DIR / "resume.txt").write_text(resume_text, encoding="utf-8")
    (DATA_DIR / "resume.json").write_text(json.dumps(resume_json, indent=2, ensure_ascii=False), encoding="utf-8")
